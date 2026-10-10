"""
Document to Speech
Upload a PDF or Word (.docx) file and download the whole thing as an MP3.

Run:  python app.py   then open http://127.0.0.1:5000
Needs an internet connection (uses Microsoft Edge's neural voices via edge-tts).
"""
import asyncio
import io
import re
import tempfile
import uuid
from pathlib import Path

import edge_tts
from docx import Document
from flask import Flask, jsonify, render_template_string, request, send_file
from pypdf import PdfReader

app = Flask(__name__)
app.config["MAX_CONTENT_LENGTH"] = 50 * 1024 * 1024  # 50 MB upload limit

OUT_DIR = Path(__file__).parent / "output"
OUT_DIR.mkdir(exist_ok=True)

VOICES = {
    "en-IN-NeerjaNeural": "English (India) - Neerja, female",
    "en-IN-PrabhatNeural": "English (India) - Prabhat, male",
    "hi-IN-SwaraNeural": "Hindi - Swara, female",
    "hi-IN-MadhurNeural": "Hindi - Madhur, male",
    "mr-IN-AarohiNeural": "Marathi - Aarohi, female",
    "mr-IN-ManoharNeural": "Marathi - Manohar, male",
    "en-US-AriaNeural": "English (US) - Aria, female",
    "en-US-GuyNeural": "English (US) - Guy, male",
    "en-GB-SoniaNeural": "English (UK) - Sonia, female",
}

CHUNK_CHARS = 2500


def extract_text(filename: str, data: bytes) -> str:
    ext = Path(filename).suffix.lower()
    if ext == ".pdf":
        reader = PdfReader(io.BytesIO(data))
        pages = [(p.extract_text() or "") for p in reader.pages]
        return "\n\n".join(pages)
    if ext == ".docx":
        doc = Document(io.BytesIO(data))
        parts = [p.text for p in doc.paragraphs if p.text.strip()]
        for table in doc.tables:
            for row in table.rows:
                cells = [c.text.strip() for c in row.cells if c.text.strip()]
                if cells:
                    parts.append(", ".join(cells))
        return "\n\n".join(parts)
    if ext == ".txt":
        return data.decode("utf-8", errors="ignore")
    raise ValueError("Unsupported file type. Please upload a PDF, DOCX or TXT file.")


def clean_text(text: str) -> str:
    text = text.replace("\u00ad", "")              # soft hyphens
    text = re.sub(r"-\n(?=[a-z])", "", text)       # join words hyphenated at line end
    text = re.sub(r"(?<!\n)\n(?!\n)", " ", text)   # single newlines -> space
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def split_chunks(text: str, limit: int = CHUNK_CHARS):
    sentences = re.split(r"(?<=[.!?।])\s+|\n\n", text)
    chunk = ""
    for s in sentences:
        s = s.strip()
        if not s:
            continue
        while len(s) > limit:  # very long sentence: hard split
            cut = s.rfind(" ", 0, limit)
            cut = cut if cut > 0 else limit
            if chunk:
                yield chunk
                chunk = ""
            yield s[:cut]
            s = s[cut:].strip()
        if len(chunk) + len(s) + 1 > limit:
            if chunk:
                yield chunk
            chunk = s
        else:
            chunk = f"{chunk} {s}".strip()
    if chunk:
        yield chunk


async def synthesize_chunk(chunk: str, voice: str, rate: str) -> bytes:
    comm = edge_tts.Communicate(chunk, voice, rate=rate)
    audio_data = bytearray()
    async for msg in comm.stream():
        if msg["type"] == "audio":
            audio_data.extend(msg["data"])
    return bytes(audio_data)


async def synthesize(text: str, voice: str, rate: str, out_path: Path):
    chunks = list(split_chunks(text))
    sem = asyncio.Semaphore(10)  # limit concurrency to avoid rate limits
    
    async def process_with_sem(chunk):
        async with sem:
            return await synthesize_chunk(chunk, voice, rate)
            
    tasks = [process_with_sem(chunk) for chunk in chunks]
    results = await asyncio.gather(*tasks)
    
    with open(out_path, "wb") as f:
        for audio_bytes in results:
            f.write(audio_bytes)


PAGE = """
<!doctype html>
<html lang="en">
<head>
<meta charset="utf-8">
<meta name="viewport" content="width=device-width, initial-scale=1">
<title>Document to Speech</title>
<style>
  :root { --bg:#f6f7f9; --card:#fff; --ink:#1d2330; --muted:#6b7385; --accent:#3b5bdb; --line:#e3e6ec; }
  @media (prefers-color-scheme: dark) {
    :root { --bg:#14161b; --card:#1d2027; --ink:#eceff4; --muted:#9aa3b5; --accent:#7b93ff; --line:#2c303a; }
  }
  * { box-sizing:border-box; }
  body { margin:0; font-family:system-ui,-apple-system,Segoe UI,Roboto,sans-serif; background:var(--bg); color:var(--ink); }
  main { max-width:560px; margin:48px auto; padding:0 16px; }
  .card { background:var(--card); border:1px solid var(--line); border-radius:14px; padding:24px; }
  h1 { font-size:22px; margin:0 0 4px; }
  p.sub { color:var(--muted); margin:0 0 20px; font-size:14px; }
  label { display:block; font-size:13px; color:var(--muted); margin:16px 0 6px; }
  select, input[type=file] { width:100%; padding:10px; border:1px solid var(--line); border-radius:8px; background:var(--bg); color:var(--ink); font-size:14px; }
  button { margin-top:22px; width:100%; padding:12px; border:0; border-radius:8px; background:var(--accent); color:#fff; font-size:15px; font-weight:600; cursor:pointer; }
  button:disabled { opacity:.6; cursor:wait; }
  #status { margin-top:16px; font-size:14px; color:var(--muted); }
  #result { margin-top:16px; display:none; }
  #result audio { width:100%; }
  #result a { display:inline-block; margin-top:10px; color:var(--accent); font-weight:600; text-decoration:none; }
  .err { color:#d6336c; }
</style>
</head>
<body>
<main>
  <div class="card">
    <h1>Document to Speech</h1>
    <p class="sub">Upload a PDF or Word file and download it as an MP3.</p>
    <form id="f">
      <label>Document (.pdf, .docx, .txt)</label>
      <input type="file" name="file" accept=".pdf,.docx,.txt" required>
      <label>Voice</label>
      <select name="voice">
        {% for k, v in voices.items() %}<option value="{{k}}">{{v}}</option>{% endfor %}
      </select>
      <label>Speed</label>
      <select name="rate">
        <option value="-20%">Slow</option>
        <option value="+0%" selected>Normal</option>
        <option value="+15%">Fast</option>
        <option value="+30%">Very fast</option>
      </select>
      <button id="go" type="submit">Convert to audio</button>
    </form>
    <div id="status"></div>
    <div id="result">
      <audio id="player" controls></audio><br>
      <div style="margin-top: 12px; display: flex; align-items: center; gap: 10px;">
        <label for="playbackSpeed" style="margin: 0; font-weight: 600;">Listen Speed:</label>
        <input type="range" id="playbackSpeed" min="0.5" max="2" step="0.1" value="1" 
               oninput="document.getElementById('player').playbackRate = this.value; document.getElementById('speedLabel').innerText = this.value + 'x';">
        <span id="speedLabel" style="font-weight: 600; color: var(--accent);">1x</span>
      </div>
      <a id="dl" download="audio.mp3">Download MP3</a>
    </div>
  </div>
</main>
<script>
const f = document.getElementById('f'), st = document.getElementById('status'),
      go = document.getElementById('go'), res = document.getElementById('result');
f.addEventListener('submit', async e => {
  e.preventDefault();
  res.style.display = 'none'; st.className = '';
  st.textContent = 'Reading the document and generating audio. Long files can take a few minutes...';
  go.disabled = true;
  try {
    const r = await fetch('/convert', { method:'POST', body:new FormData(f) });
    let j;
    const text_res = await r.text();
    try {
      j = JSON.parse(text_res);
    } catch (parseErr) {
      if (!r.ok) throw new Error(`Server returned an error (${r.status}). This usually means the file took too long to process and the request timed out.`);
      throw new Error('Invalid response from server.');
    }
    if (!r.ok) throw new Error(j.error || 'Something went wrong');
    document.getElementById('player').src = j.url;
    const dl = document.getElementById('dl');
    dl.href = j.url + '?dl=1'; dl.download = j.name;
    document.getElementById('playbackSpeed').value = 1;
    document.getElementById('speedLabel').innerText = '1x';
    res.style.display = 'block';
    st.textContent = `Done. ${j.chars.toLocaleString()} characters converted.`;
  } catch (err) {
    st.className = 'err'; st.textContent = err.message;
  } finally { go.disabled = false; }
});
</script>
</body>
</html>
"""


@app.get("/")
def index():
    return render_template_string(PAGE, voices=VOICES)


@app.post("/convert")
def convert():
    file = request.files.get("file")
    if not file or not file.filename:
        return jsonify(error="Please choose a file."), 400
    voice = request.form.get("voice", "en-IN-NeerjaNeural")
    rate = request.form.get("rate", "+0%")
    if voice not in VOICES:
        return jsonify(error="Unknown voice."), 400
    try:
        text = clean_text(extract_text(file.filename, file.read()))
    except ValueError as e:
        return jsonify(error=str(e)), 400
    except Exception:
        return jsonify(error="Could not read that file. Is it corrupted or password protected?"), 400
    if not text:
        return jsonify(error="No readable text found. Scanned PDFs (images) need OCR first."), 400

    OUT_DIR.mkdir(exist_ok=True)
    out_id = uuid.uuid4().hex
    out_path = OUT_DIR / f"{out_id}.mp3"
    try:
        asyncio.run(synthesize(text, voice, rate, out_path))
    except Exception as e:
        return jsonify(error=f"Speech generation failed (check your internet connection): {e}"), 500

    name = Path(file.filename).stem + ".mp3"
    (OUT_DIR / f"{out_id}.name").write_text(name, encoding="utf-8")
    return jsonify(url=f"/audio/{out_id}", name=name, chars=len(text))


@app.get("/audio/<out_id>")
def audio(out_id):
    if not re.fullmatch(r"[0-9a-f]{32}", out_id):
        return "Not found", 404
    path = OUT_DIR / f"{out_id}.mp3"
    if not path.exists():
        return "Not found", 404
    name_file = OUT_DIR / f"{out_id}.name"
    name = name_file.read_text(encoding="utf-8") if name_file.exists() else "audio.mp3"
    return send_file(path, mimetype="audio/mpeg",
                     as_attachment=bool(request.args.get("dl")), download_name=name)


if __name__ == "__main__":
    app.run(host="127.0.0.1", port=5000, debug=False)
